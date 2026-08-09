"""
Section 4 (Support Plan) template engine.

`Templates/Section_four/Support Plan/` contains one Word document per
support sub-area, laid out as:

    Support Plan/
      <Area covered>/
        <Sub-area covered>.docx
      ...

Every one of those sub-area documents shares the same fundamental
structure (verified by inspection): a heading with the area name, a heading
with the sub-area name, and a single table with columns
`Objective Steps | Date Completed | Client Signature | Support Worker
Signature`, a variable number of "Step N: ..." rows, and a trailing
"Comments:" row. Only the number of step rows differs between files.

This module turns that folder layout + structure into two reusable pieces:

- `SupportPlanRegistry`: an index of every sub-area template file, so any
  caller can look up the exact template for an (area_covered,
  sub_area_covered) pair - e.g. one entry from a JSON
  `list_of_support_areas_covered` array - without needing to know the
  folder layout, and without ever guessing at file paths.
- Step-table helpers (`list_steps`, `complete_step`, `set_comments`): a
  generic way to read and fill in the Objective Steps table of a template
  once it has been matched and opened, regardless of how many steps that
  particular sub-area happens to have.

Nothing in this module assumes a specific input JSON schema - it only
knows how to match strings to a template file and how to write into that
file's known table structure. That keeps it reusable by whatever Section 4
generator/JSON schema is designed later.
"""

from __future__ import annotations

import difflib
import random
import re
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import List, Mapping, Optional, Sequence, Tuple, Union

from docx import Document
from docx.table import Table

DEFAULT_SUPPORT_PLAN_DIR = (
    Path(__file__).parent.parent / "Templates" / "Section_four" / "Support Plan"
)

# The Objective Steps table's fixed header, used both to locate that table
# amongst any others in a document and to know its column order.
STEP_TABLE_HEADER = ["Objective Steps", "Date Completed", "Client Signature", "Support Worker Signature"]
DATE_COMPLETED_COLUMN = 1
CLIENT_SIGNATURE_COLUMN = 2
SUPPORT_WORKER_SIGNATURE_COLUMN = 3

# Step labels are not always plain integers - some templates split a step
# into lettered sub-steps (e.g. "Step 1a:", "Step 1b:"). The identifier is
# therefore kept as the exact string that follows "Step " (e.g. "1", "1a"),
# rather than assumed to be an int.
_STEP_ROW_PATTERN = re.compile(r"^Step\s+(\w+)\s*:\s*(.*)$", re.IGNORECASE | re.DOTALL)
_COMMENTS_ROW_LABEL = "Comments:"


def _is_word_lock_file(path: Path) -> bool:
    """Word creates hidden `~$<name>.docx` lock files while a doc is open."""
    return path.name.startswith("~$")


# --------------------------------------------------------------------------
# Template lookup
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class SupportAreaTemplate:
    """
    One Section 4 support-plan template file, addressed by the exact
    (area_covered, sub_area_covered) pair it was filed under.
    """

    area_covered: str
    sub_area_covered: str
    template_path: Path

    def open(self) -> Document:
        """Open a fresh, independent copy of this template ready for editing."""
        return Document(self.template_path)

    def copy_to(self, destination_path: Union[str, Path], *, overwrite: bool = False) -> Path:
        """
        Copy this template file to `destination_path`, creating any missing
        parent folders. If a file already exists there, it is left
        untouched unless `overwrite=True` - this lets callers safely call
        `copy_to` once per session without wiping out a document that
        earlier sessions have already started writing into.
        """
        destination_path = Path(destination_path)
        if destination_path.exists() and not overwrite:
            return destination_path
        destination_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(self.template_path, destination_path)
        return destination_path


def _extract_area_strings(source: Union[Mapping[str, str], object]) -> Tuple[str, str]:
    """
    Read (area_covered, sub_area_covered) off of `source`, which may be a
    dict-like object (e.g. a JSON-decoded entry) or any object exposing
    those two fields as attributes.
    """
    if isinstance(source, Mapping):
        try:
            return source["area_covered"], source["sub_area_covered"]
        except KeyError as exc:
            raise ValueError(f"{source!r} is missing required key {exc}.") from exc

    try:
        return source.area_covered, source.sub_area_covered
    except AttributeError as exc:
        raise ValueError(
            f"{source!r} must be a mapping with 'area_covered'/'sub_area_covered' keys, "
            "or an object exposing them as attributes."
        ) from exc


class SupportPlanRegistry:
    """
    Indexes every Section 4 support-plan template file under a
    `Support Plan/<Area covered>/<Sub-area covered>.docx` folder tree, so
    callers can look up the exact template for any (area_covered,
    sub_area_covered) pair by string comparison alone.
    """

    def __init__(self, support_plan_dir: Union[str, Path] = DEFAULT_SUPPORT_PLAN_DIR):
        self.support_plan_dir = Path(support_plan_dir)
        self._templates: List[SupportAreaTemplate] = self._scan(self.support_plan_dir)

    @staticmethod
    def _scan(support_plan_dir: Path) -> List[SupportAreaTemplate]:
        if not support_plan_dir.is_dir():
            raise FileNotFoundError(f"Support Plan template folder not found: {support_plan_dir}")

        templates = []
        for area_folder in sorted(support_plan_dir.iterdir()):
            if not area_folder.is_dir():
                continue  # Skip top-level files (cover page, blank master template).
            area_covered = area_folder.name
            for docx_path in sorted(area_folder.glob("*.docx")):
                if _is_word_lock_file(docx_path):
                    continue
                sub_area_covered = docx_path.stem
                templates.append(SupportAreaTemplate(area_covered, sub_area_covered, docx_path))
        return templates

    def all(self) -> Sequence[SupportAreaTemplate]:
        """Every indexed template, in folder order."""
        return tuple(self._templates)

    def areas(self) -> List[str]:
        """All distinct main support areas, in folder order."""
        seen: List[str] = []
        for template in self._templates:
            if template.area_covered not in seen:
                seen.append(template.area_covered)
        return seen

    def sub_areas(self, area_covered: str) -> List[str]:
        """All sub-areas filed under one main area."""
        return [t.sub_area_covered for t in self._templates if t.area_covered == area_covered]

    def find(self, area_covered: str, sub_area_covered: str) -> SupportAreaTemplate:
        """
        Return the template whose area/sub-area match EXACTLY (case- and
        spelling-sensitive, per the support-plan spec - wording must never
        be altered). Raises ValueError listing the closest known values if
        nothing matches exactly.
        """
        for template in self._templates:
            if template.area_covered == area_covered and template.sub_area_covered == sub_area_covered:
                return template
        raise ValueError(self._not_found_message(area_covered, sub_area_covered))

    def match(self, source: Union[Mapping[str, str], object]) -> SupportAreaTemplate:
        """
        Same as `find`, but reads area_covered/sub_area_covered directly off
        of `source` - typically one entry from a JSON
        `list_of_support_areas_covered` array, or a session object that
        exposes those two fields.
        """
        area_covered, sub_area_covered = _extract_area_strings(source)
        return self.find(area_covered, sub_area_covered)

    def _not_found_message(self, area_covered: str, sub_area_covered: str) -> str:
        known_areas = self.areas()
        if area_covered not in known_areas:
            closest = difflib.get_close_matches(area_covered, known_areas, n=3)
            hint = f" Closest known areas: {closest}." if closest else ""
            return f'Unknown area_covered "{area_covered}".{hint}'

        known_sub_areas = self.sub_areas(area_covered)
        closest = difflib.get_close_matches(sub_area_covered, known_sub_areas, n=3)
        hint = f' Closest known sub-areas for "{area_covered}": {closest}.' if closest else ""
        return f'Unknown sub_area_covered "{sub_area_covered}" under "{area_covered}".{hint}'


# --------------------------------------------------------------------------
# Objective Steps table helpers
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class StepRow:
    """
    One "Step <identifier>: ..." row found in a template's Objective Steps
    table. `identifier` is kept as the exact string following "Step "
    (e.g. "1", "2", "1a"), since some templates split a step into lettered
    sub-steps rather than using plain integers throughout.
    """

    identifier: str
    description: str
    row_index: int


def get_steps_table(document: Document) -> Table:
    """Return the Objective Steps table, identified by its fixed header row."""
    for table in document.tables:
        header_texts = [cell.text.strip() for cell in table.rows[0].cells]
        if header_texts == STEP_TABLE_HEADER:
            return table
    raise ValueError("Could not find the Objective Steps table in this document.")


def list_steps(document: Document) -> List[StepRow]:
    """
    Return every "Step <identifier>: ..." row from the Objective Steps
    table, in document order (skips the header row and the trailing
    Comments row automatically, since neither matches the "Step ...:"
    pattern).
    """
    table = get_steps_table(document)
    steps = []
    for row_index, row in enumerate(table.rows):
        match = _STEP_ROW_PATTERN.match(row.cells[0].text.strip())
        if match:
            steps.append(StepRow(identifier=match.group(1), description=match.group(2).strip(), row_index=row_index))
    return steps


def _set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    for run in list(paragraph.runs):
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def complete_step(
    document: Document,
    step_identifier: Union[str, int],
    *,
    date_completed: Optional[str] = None,
    client_signature: Optional[str] = None,
    support_worker_signature: Optional[str] = None,
) -> None:
    """
    Fill in the Date Completed / Client Signature / Support Worker Signature
    cells for one step row, identified by its exact label (e.g. "1", "2",
    or "1a" for a lettered sub-step - ints are accepted and compared as
    strings for convenience). Any argument left as None is not written,
    leaving that cell untouched.
    """
    step_identifier = str(step_identifier)
    table = get_steps_table(document)
    steps_by_identifier = {step.identifier: step for step in list_steps(document)}
    if step_identifier not in steps_by_identifier:
        raise ValueError(f'Step "{step_identifier}" not found. Available steps: {sorted(steps_by_identifier)}')

    row = table.rows[steps_by_identifier[step_identifier].row_index]
    if date_completed is not None:
        _set_cell_text(row.cells[DATE_COMPLETED_COLUMN], date_completed)
    if client_signature is not None:
        _set_cell_text(row.cells[CLIENT_SIGNATURE_COLUMN], client_signature)
    if support_worker_signature is not None:
        _set_cell_text(row.cells[SUPPORT_WORKER_SIGNATURE_COLUMN], support_worker_signature)


def set_comments(
    document: Document,
    *,
    objective_steps: Optional[str] = None,
    date_completed: Optional[str] = None,
    client_signature: Optional[str] = None,
    support_worker_signature: Optional[str] = None,
) -> None:
    """
    Fill in the trailing "Comments:" row, one column at a time, keeping each
    cell's "Comments:" label and appending the supplied text after it. Any
    argument left as None leaves that column untouched.
    """
    table = get_steps_table(document)
    comments_row = table.rows[-1]
    values = [objective_steps, date_completed, client_signature, support_worker_signature]
    for column_index, value in enumerate(values):
        if value is not None:
            _set_cell_text(comments_row.cells[column_index], f"{_COMMENTS_ROW_LABEL} {value}")


def append_comment_entry(
    document: Document,
    entry_text: str,
    *,
    column: int = 0,
    add_blank_line_before: Optional[bool] = None,
) -> None:
    """
    Append one line of free text to a Comments row cell, without disturbing
    whatever is already there (the template's own blank space, or earlier
    entries from previous calls) - unlike `set_comments`, which overwrites a
    cell outright. This is meant for logging one entry per occurrence (e.g.
    one per support session that covers this sub-area), building up a
    running list of entries in a single cell over multiple calls.

    `column` selects which of the four Comments cells to append to
    (0=Objective Steps, 1=Date Completed, 2=Client Signature,
    3=Support Worker Signature) - it defaults to the Objective Steps
    column, the main free-text column.

    Consecutive entries are separated by either 0 or 1 blank line: if
    `add_blank_line_before` is left as None, this is decided randomly
    (50/50) on each call; pass True/False for deterministic control (e.g.
    in tests).
    """
    comments_row = get_steps_table(document).rows[-1]
    cell = comments_row.cells[column]
    if add_blank_line_before is None:
        add_blank_line_before = random.random() < 0.5
    if add_blank_line_before:
        cell.add_paragraph()
    cell.add_paragraph(entry_text)
